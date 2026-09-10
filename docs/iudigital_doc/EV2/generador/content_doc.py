# -*- coding: utf-8 -*-
"""Contenido del documento de la Actividad 2 de CreditSmart (interfaz en React)."""
from build_doc import *   # noqa: F401,F403  (helpers, estilos y objeto `doc`)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

configure_styles()
build_header_footer()

C = WD_ALIGN_PARAGRAPH.CENTER
J = WD_ALIGN_PARAGRAPH.JUSTIFY

# ============================================================== PORTADA
for text in (u'S30 · Actividad 2 — Evidencia de Aprendizaje 2',
             u'Ingeniería Web',
             u'CreditSmart — Aplicación Web Dinámica con React'):
    p = doc.add_paragraph(text, style='Heading 1')
    p.alignment = C
    p.paragraph_format.space_after = Pt(2)

para(u'Documento técnico: migración de la interfaz a React, componentes y props, '
     u'manejo de estado con hooks, búsqueda y filtros dinámicos, formulario '
     u'controlado y cálculo de la cuota mensual, sobre la arquitectura hexagonal '
     u'construida en la Actividad 1.',
     align=C, size=11, italic=True, color='6B7280', space_after=18)

para(u'Integrantes', align=C, size=11, bold=True, space_after=2)
para(u'Jeremy Ivan Pedraza Hernández', align=C, size=11, space_after=1)
para(u'Jossy Esteban Paternina Julio', align=C, size=11, space_after=8)

para(u'Docente', align=C, size=11, bold=True, space_after=2)
para(u'Jorge Armando Julio', align=C, size=11, space_after=8)

para(u'Curso y NRC', align=C, size=11, bold=True, space_after=2)
para(u'Ingeniería Web — PREICA2602B010133', align=C, size=11, space_after=14)

para(u'Ingeniería de Software y Datos', align=C, bold=True, space_after=2)
para(u'Facultad de Ingeniería y Ciencias Agropecuarias', align=C, space_after=2)
para(u'Institución Universitaria Digital de Antioquia', align=C, space_after=2)
para(u'2026-S2', align=C, space_after=10)
field_placeholder(u'[ Ciudad y fecha de entrega ]', width_cm=8.0)

page_break()

# ============================================================== TABLA DE CONTENIDO
h1(u'Tabla de contenido')
para(u'Para actualizar el índice: seleccione la tabla, haga clic derecho y elija '
     u'«Actualizar campos» → «Actualizar toda la tabla».', size=9,
     italic=True, color='6B7280')
add_toc()
page_break()

# ============================================================== 1. INTRODUCCION
h1(u'1. Introducción')
para(u'CreditSmart es una plataforma web para consultar, simular y solicitar '
     u'productos de crédito. En la Actividad 1 se construyó con HTML5, CSS3 y '
     u'JavaScript estándar, sin frameworks y sin proceso de compilación, sobre una '
     u'arquitectura hexagonal con la regla de dependencia de Clean Architecture y '
     u'los cinco principios SOLID.')
para(u'Esta Actividad 2 transforma ese diseño estático en una aplicación web '
     u'interactiva con React: componentes reutilizables con props, manejo de estado '
     u'con hooks, enrutado con React Router, búsqueda en tiempo real, filtros y '
     u'ordenamiento dinámicos, formulario completamente controlado con validaciones '
     u'y cálculo automático de la cuota mensual.')
para(u'La tesis del documento es que esa transformación no fue una reescritura, sino '
     u'la sustitución de un adaptador. La arquitectura hexagonal separa la interfaz '
     u'de las reglas del negocio, de modo que cambiar de tecnología de interfaz solo '
     u'obliga a reescribir la interfaz. El resultado, medido en archivos: la capa de '
     u'dominio no cambió en absoluto; la de aplicación creció en un caso de uso; la '
     u'de presentación se reescribió por completo.')

callout(u'Cómo leer este documento',
        u'Las figuras resumen visualmente cada concepto y las tablas contienen el '
        u'detalle exacto (nombres de archivo, firmas, valores). Los bloques grises '
        u'son fragmentos reales del código fuente. Los recuadros amarillos son '
        u'espacios editables que hay que completar antes de entregar.')

# ============================================================== 2. OBJETIVO
h1(u'2. Objetivo')
h2(u'2.1 Objetivo general')
para(u'Aplicar los conceptos de React JS, la programación orientada a componentes, el '
     u'manejo de estado con hooks y la manipulación dinámica de datos para '
     u'transformar el diseño estático de CreditSmart en una aplicación web '
     u'interactiva y funcional, implementando búsquedas en tiempo real, filtros '
     u'dinámicos y formularios controlados, sin degradar la separación entre la '
     u'interfaz y las reglas del negocio.')

h2(u'2.2 Objetivos específicos')
for t in (u'Configurar un proyecto React con Vite y React Router, con una estructura '
          u'de carpetas organizada (components/, pages/, data/, hooks/).',
          u'Diseñar componentes funcionales reutilizables que reciban sus datos por '
          u'props desestructuradas, uno por archivo y sin estado propio.',
          u'Gestionar el estado de la interfaz con useState, useEffect, useMemo, '
          u'useCallback y useContext, y encapsularlo en hooks propios.',
          u'Implementar búsqueda incremental mientras el usuario escribe, filtro por '
          u'rango de monto y ordenamiento, con .map(), .filter() y .sort().',
          u'Construir un formulario 100 % controlado con validación en tiempo real de '
          u'correo, cédula, montos e ingresos, reutilizando las reglas del dominio.',
          u'Calcular la cuota mensual de forma automática con la tasa del producto '
          u'elegido y presentarla en formato COP, junto con la tabla de amortización.',
          u'Demostrar, con evidencia medible, que la arquitectura hexagonal hizo de '
          u'la migración un cambio de adaptador y no una reescritura.'):
    bullet(t)

# ============================================================== 3. ALCANCE
h1(u'3. Alcance y requerimientos')
h2(u'3.1 Requerimientos funcionales')
add_table(
    [u'Cód.', u'Requerimiento', u'Dónde se implementa'],
    [[u'RF-01', u'Mostrar el catálogo de 6 productos con nombre, descripción, tasa, '
                u'plazo, rango de montos y requisitos.',
      u'CatalogPage + CreditCard + useCreditProducts + ListCreditProductsUseCase'],
     [u'RF-02', u'Buscar productos por nombre o descripción, filtrando mientras el '
                u'usuario escribe, sin pulsar ningún botón.',
      u'SearchBar + useCreditSearch + SearchCreditProductsUseCase'],
     [u'RF-03', u'Filtrar por rango de monto (5 rangos) y limpiar los filtros.',
      u'AmountRangeFilter + GetAmountRangeFiltersUseCase'],
     [u'RF-04', u'Ordenar el catálogo por nombre, tasa, monto máximo o plazo.',
      u'SortSelect + useProductSorting'],
     [u'RF-05', u'Calcular la cuota mensual, el total de intereses, el total a pagar '
                u'y el coste del crédito al cambiar producto, monto o plazo.',
      u'SimulatorForm + useSimulation + SimulateCreditUseCase + '
      u'CreditSimulationService'],
     [u'RF-06', u'Mostrar la tabla de amortización en dos niveles de detalle: '
                u'resumen anual y mes a mes.',
      u'AmortizationTable + AmortizationPlan.yearlySummary()'],
     [u'RF-07', u'Capturar la solicitud de crédito: 11 campos en 3 secciones, con el '
                u'formulario 100 % controlado.',
      u'ApplicationPage + FormField + useApplicationForm'],
     [u'RF-08', u'Validar en tiempo real correo, cédula, teléfono, montos, plazo e '
                u'ingresos, con mensajes claros por campo.',
      u'ValidateCreditApplicationDraftUseCase + Applicant · RequestedCredit · '
      u'EmploymentInfo'],
     [u'RF-09', u'Radicar la solicitud, avisar al usuario con el número de radicado y '
                u'limpiar el formulario.',
      u'SubmitCreditApplicationUseCase + LocalStorageCreditApplicationRepository + '
      u'ToastNotifier'],
     [u'RF-10', u'Navegar entre catálogo, simulador y solicitud sin recargar la '
                u'página, con una pantalla 404 para rutas inexistentes.',
      u'App.jsx + React Router (BrowserRouter, Routes, NavLink, Link)']],
    widths=[1.5, 7.4, 7.5])
cap(u'Requerimientos funcionales y su implementación', kind='Tabla')

h2(u'3.2 Requerimientos no funcionales')
add_table(
    [u'Cód.', u'Requerimiento', u'Cómo se cumple'],
    [[u'RNF-01', u'La interfaz responde a cada pulsación sin recargar y sin esperas '
                 u'perceptibles.',
      u'Estado local con hooks; el catálogo se resuelve en memoria'],
     [u'RNF-02', u'Diseño responsive mobile-first verificado en móvil, tablet y '
                 u'escritorio.',
      u'Los 7 CSS de la Actividad 1, reutilizados sin cambios'],
     [u'RNF-03', u'Las reglas de negocio no dependen de React ni del navegador.',
      u'domain/ y application/ son módulos ES estándar; la suite corre en Node'],
     [u'RNF-04', u'Un error de configuración debe fallar al arrancar, no a mitad de '
                 u'una navegación.',
      u'assertImplements + container.eagerResolveAll() en main.jsx'],
     [u'RNF-05', u'Accesibilidad: etiquetas asociadas, errores anunciados, foco '
                 u'visible.',
      u'label htmlFor, role="alert", aria-invalid, aria-live en los avisos'],
     [u'RNF-06', u'Seguridad del marcado: ningún dato del usuario se interpreta como '
                 u'HTML.',
      u'JSX escapa por defecto; no se usa dangerouslySetInnerHTML en ningún archivo'],
     [u'RNF-07', u'El proyecto se instala y arranca con dos comandos.',
      u'npm install · npm run dev (Vite)']],
    widths=[1.7, 7.3, 7.4])
cap(u'Requerimientos no funcionales', kind='Tabla')

page_break()

# ============================================================== 4. VISION GENERAL
h1(u'4. Visión general de la solución')
h2(u'4.1 Cifras de la solución')
add_table(
    [u'Dato', u'Valor'],
    [[u'Stack', u'React 19 · React Router 7 · Vite 8 · CSS3 plano · JavaScript ES2022'],
     [u'Archivos JS/JSX', u'82 en total'],
     [u'Capa de dominio', u'26 archivos — 2 entidades, 10 value objects, 2 servicios, '
      u'6 puertos, 4 errores, 1 criterio, 1 fábrica de contratos'],
     [u'Capa de aplicación', u'16 archivos — 7 casos de uso, 2 DTOs, 3 mappers, '
      u'3 puertos, Result'],
     [u'Capa de infraestructura', u'10 adaptadores'],
     [u'Capa de presentación (React)', u'25 archivos — 11 componentes, 7 hooks, '
      u'4 páginas, 1 contexto, App.jsx, main.jsx'],
     [u'Configuración y datos', u'4 + 1 archivos — AppConfig, Container, '
      u'dependencies, routes; creditsData'],
     [u'Archivos CSS', u'7 (1 681 líneas) reutilizados sin cambios de la Actividad 1'],
     [u'Contratos (puertos)', u'9 — 6 de dominio y 3 de aplicación'],
     [u'Dependencias en el contenedor', u'20'],
     [u'Productos del catálogo', u'6'],
     [u'Rutas', u'/ · /simulador · /solicitar · * (404)'],
     [u'Hooks propios', u'7'],
     [u'Pruebas', u'tests/01-domain-application.mjs — 85 aserciones, TODO OK'],
     [u'Commits de la entrega', u'9 en la rama ev2-react (+1 en main)']],
    widths=[5.0, 11.4])
cap(u'Cifras de la solución', kind='Tabla')

h2(u'4.2 Responsabilidad de cada capa')
add_table(
    [u'Capa', u'Responsabilidad', u'Puede', u'No puede'],
    [[u'Dominio', u'Reglas e invariantes del negocio',
      u'Definir puertos, lanzar errores de negocio',
      u'Conocer React, el DOM, la red o el almacenamiento'],
     [u'Aplicación', u'Orquestar casos de uso',
      u'Llamar al dominio y a los puertos, devolver Result con DTOs',
      u'Tener reglas propias o tocar el DOM'],
     [u'Infraestructura', u'Implementar los puertos con tecnología concreta',
      u'Usar localStorage, Intl, crypto, Date',
      u'Contener reglas de negocio'],
     [u'Presentación (React)', u'Pintar y capturar la intención del usuario',
      u'Guardar estado de interfaz e invocar casos de uso',
      u'Importar infraestructura, entidades o value objects'],
     [u'Configuración', u'Construir el grafo de dependencias',
      u'Hacer new de clases concretas',
      u'Ser importada por dominio, aplicación o componentes']],
    widths=[2.9, 3.9, 4.9, 4.7])
cap(u'Responsabilidad de cada capa', kind='Tabla')

boxes_row([(u'PRESENTACIÓN\nReact', BRAND), (u'APLICACIÓN\ncasos de uso', VIOLET),
           (u'DOMINIO\nreglas', NAVY), (u'INFRAESTRUCTURA\nadaptadores', AMBER)])
cap(u'Flujo de control: la interfaz pide, la aplicación orquesta, el dominio decide '
    u'y la infraestructura ejecuta. Las dependencias, en cambio, apuntan siempre '
    u'hacia el dominio.')

page_break()

# ============================================================== 5. DE EV1 A EV2
h1(u'5. De la Actividad 1 a la Actividad 2')
h2(u'5.1 Qué cambió, medido en archivos')
para(u'La Actividad 1 exigía JavaScript sin frameworks; la Actividad 2 exige React. '
     u'Si la interfaz y las reglas de negocio hubieran estado mezcladas, cambiar de '
     u'tecnología habría significado reescribir la aplicación completa. No fue el '
     u'caso.')
figure('chart-migracion.png',
       u'Archivos por capa antes y después de la migración. El dominio queda idéntico; '
       u'solo la presentación se reescribe.', width_cm=15.5)

add_table(
    [u'Capa', u'Antes', u'Después', u'Cambio'],
    [[u'domain/', u'26', u'26', u'Ninguno'],
     [u'application/', u'15', u'16', u'1 caso de uso nuevo · 1 modificado'],
     [u'infrastructure/', u'11', u'10', u'−1: HistoryRouter, que ahora hace React Router'],
     [u'config/', u'4', u'4', u'dependencies.js deja de registrar la presentación'],
     [u'data/', u'—', u'1', u'Catálogo extraído a creditsData.js'],
     [u'presentación', u'22', u'25', u'Reescrita en React']],
    widths=[3.6, 2.0, 2.2, 8.6], align_center_cols=(1, 2))
cap(u'Impacto de la migración por capa', kind='Tabla')

h2(u'5.2 Tabla de equivalencias')
para(u'Ninguna pieza se perdió: cambió su forma de expresarse.')
add_table(
    [u'Actividad 1 (vanilla)', u'Actividad 2 (React)', u'Qué hace'],
    [[u'presentation/views/*View.js', u'src/pages/*Page.jsx', u'Pintar una pantalla'],
     [u'presentation/controllers/*Controller.js', u'src/hooks/use*.js',
      u'Guardar el estado de interfaz e invocar casos de uso'],
     [u'presentation/components/*Component.js', u'src/components/*.jsx',
      u'Componentes reutilizables sin estado'],
     [u'BaseView + this.on(...)', u'React', u'Ciclo de vida y limpieza de listeners'],
     [u'ViewRenderer (innerHTML)', u'React DOM', u'Escribir en el DOM'],
     [u'Html.js (html, raw, escapeHtml)', u'JSX', u'Escapado por defecto'],
     [u'HistoryRouter + UrlBuilder',
      u'BrowserRouter · Routes · NavLink · Link', u'Navegación sin recarga'],
     [u'DocumentTitleController (decorador)', u'useDocumentTitle',
      u'Título del documento por ruta'],
     [u'container.resolve(...) en main.js',
      u'DependenciesProvider + useDependencies',
      u'Entregar las dependencias a la interfaz'],
     [u'IView · IController · IRouter', u'—',
      u'Contratos que React ya impone por la firma de un componente o un hook']],
    widths=[5.4, 5.0, 6.0])
cap(u'Equivalencias entre la presentación vanilla y la de React', kind='Tabla')

para(u'Los tres contratos de presentación desaparecen porque su función era forzar '
     u'una forma que en React viene dada: un componente es una función que devuelve '
     u'JSX y un hook devuelve estado. Los nueve puertos de dominio y aplicación '
     u'siguen intactos, porque esos no describen la interfaz: describen el negocio.')

h2(u'5.3 Los tres cambios que sí exigió React')
h3(u'5.3.1 La búsqueda recibe un índice, no un value object')
para(u'En la Actividad 1 el controlador del simulador recibía inyectado el puerto '
     u'IAmountRangeProvider, resolvía el AmountRange y lo pasaba al caso de uso. Un '
     u'componente de React no debe construir value objects del dominio, así que el '
     u'caso de uso pasó a aceptar un índice y a resolver el rango contra el puerto.')
code_block([
    u'// application/usecases/SearchCreditProductsUseCase.js',
    u'async execute({ query = \'\', amountRange = null, amountRangeIndex = null } = {}) {',
    u'  const range = amountRange ?? (await this.#resolveRange(amountRangeIndex));',
    u'  const criteria = new ProductSearchCriteria({ query, amountRange: range });',
    u'  const products = await this.#repository.findByCriteria(criteria);',
    u'  return Result.ok({ products: this.#mapper.toDTOList(products), /* … */ });',
    u'}',
], caption=u'La interfaz manda primitivos y recibe DTOs: nada más cruza la frontera.')

h3(u'5.3.2 Un caso de uso para validar sin duplicar reglas')
para(u'La rúbrica pide validación en tiempo real de correo, cédula y montos. La '
     u'tentación es escribir esas validaciones en el componente, y entonces existen '
     u'dos verdades sobre lo que es válido: la del formulario y la del dominio. El '
     u'caso de uso nuevo valida un borrador reutilizando los value objects, sin '
     u'persistir nada.')
code_block([
    u'// application/usecases/ValidateCreditApplicationDraftUseCase.js',
    u'async execute(rawForm = {}) {',
    u'  try {',
    u'    CreditApplicationMapper.toValueObjects(rawForm);   // Applicant,',
    u'    return Result.ok({ isValid: true });               // RequestedCredit,',
    u'  } catch (err) {                                      // EmploymentInfo',
    u'    return Result.fromError(err);   // fieldErrors planos, listos para pintar',
    u'  }',
    u'}',
], caption=u'El mensaje que ve el usuario al escribir es el que produjo el dominio.')

h3(u'5.3.3 El catálogo pasa a src/data/creditsData.js')
para(u'La rúbrica pide una carpeta data/ con el archivo de datos. Los arrays crudos '
     u'se movieron allí y StaticCreditProductDataSource los reexporta con los nombres '
     u'que ya usaban los repositorios. No rompe la regla de dependencia porque '
     u'creditsData.js no importa nada: es dato puro, sin capa.')

page_break()

# ============================================================== 6. CONFIGURACION
h1(u'6. Configuración y estructura del proyecto')
h2(u'6.1 Herramientas y versiones')
add_table(
    [u'Tecnología', u'Versión', u'Para qué'],
    [[u'React', u'19', u'Componentes, estado con hooks, render declarativo'],
     [u'React Router', u'7', u'Enrutado SPA de las cuatro rutas'],
     [u'Vite', u'8', u'Servidor de desarrollo con recarga en caliente y empaquetado'],
     [u'@vitejs/plugin-react', u'6', u'Transformación de JSX'],
     [u'Node.js', u'22.14', u'Entorno de desarrollo y ejecución de la suite'],
     [u'CSS3', u'—', u'7 hojas en cascada, Grid y Flexbox, mobile-first']],
    widths=[4.2, 2.2, 10.0], align_center_cols=(1,))
cap(u'Tecnologías utilizadas', kind='Tabla')

para(u'Se eligió Vite sobre Create React App por ser la opción recomendada en el '
     u'enunciado, por su arranque casi instantáneo y porque su salida es HTML, CSS y '
     u'JavaScript estáticos, servibles por el mismo Apache de Laragon que ya servía '
     u'la Actividad 1.')

h2(u'6.2 Comandos')
code_block([
    u'npm install        # instala React, React Router y Vite',
    u'npm run dev        # servidor de desarrollo -> http://localhost:5173',
    u'npm run build      # empaqueta para produccion en dist/',
    u'npm run preview    # sirve dist/ para revisar el empaquetado',
    u'',
    u'node tests/01-domain-application.mjs   # 85 aserciones sobre dominio y aplicacion',
], caption=u'Instalación, ejecución y verificación.')

h2(u'6.3 Estructura de carpetas')
code_block([
    u'crediSmart/',
    u'|- index.html                  Punto de entrada de Vite',
    u'|- package.json                Dependencias y scripts',
    u'|- vite.config.js              Configuracion del empaquetador',
    u'|- public/.htaccess            Reescritura SPA para Apache (se copia a dist/)',
    u'|- assets/css/                 7 hojas en cascada (reutilizadas de la Actividad 1)',
    u'|- src/',
    u'|  |- main.jsx                 Composition Root: construye el grafo y monta React',
    u'|  |- App.jsx                  Tabla de rutas (React Router)',
    u'|  |- data/creditsData.js      Catalogo, rangos de monto y plazos (dato puro)',
    u'|  |- components/              11 componentes reutilizables',
    u'|  |- pages/                   4 paginas, una por ruta',
    u'|  |- hooks/                   7 hooks: estado de UI + invocacion de casos de uso',
    u'|  |- context/                 DependenciesProvider',
    u'|  |- domain/                  Entidades, value objects, servicios, puertos',
    u'|  |- application/             Casos de uso, DTOs, mappers, Result',
    u'|  |- infrastructure/          Adaptadores: persistencia, formato, reloj, ids',
    u'|  \\- config/                  AppConfig, Container, dependencies, routes',
    u'|- tests/01-domain-application.mjs',
    u'\\- docs/                       21 documentos de diseno + capturas',
], caption=u'Estructura del proyecto. components/, pages/ y data/ están en la raíz de '
           u'src/, como pide el enunciado.')

figure('chart-react.png', u'Composición de la capa de presentación en React.',
       width_cm=13.0)

callout(u'Las carpetas de la rúbrica y las capas de la arquitectura no compiten',
        u'components/, pages/, hooks/ y context/ SON la capa de presentación. Las '
        u'otras tres capas viven en sus propias carpetas dentro de src/, y ninguna '
        u'de ellas importa a la presentación.')

page_break()

# ============================================================== 7. ENRUTADO
h1(u'7. Enrutado con React Router')
para(u'Las cuatro rutas son las mismas de la Actividad 1 y se siguen leyendo de '
     u'config/routes.js, un archivo que no importa nada: así cualquier componente '
     u'puede construir un enlace sin crear dependencias circulares con las páginas.')
add_table(
    [u'Ruta', u'Página', u'Título del documento'],
    [[u'/', u'CatalogPage', u'CreditSmart — Catálogo'],
     [u'/simulador', u'SimulatorPage', u'CreditSmart — Simulador'],
     [u'/solicitar', u'ApplicationPage', u'CreditSmart — Solicitar'],
     [u'*', u'NotFoundPage', u'CreditSmart — Página no encontrada']],
    widths=[3.4, 4.6, 8.4])
cap(u'Rutas expuestas al usuario', kind='Tabla')

code_block([
    u'// src/App.jsx',
    u'export function App() {',
    u'  return (',
    u'    <Routes>',
    u'      <Route path={ROUTES.CATALOG}     element={<CatalogPage />} />',
    u'      <Route path={ROUTES.SIMULATOR}   element={<SimulatorPage />} />',
    u'      <Route path={ROUTES.APPLICATION} element={<ApplicationPage />} />',
    u'      <Route path="*"                  element={<NotFoundPage />} />',
    u'    </Routes>',
    u'  );',
    u'}',
], caption=u'Tabla de rutas declarativa.')

para(u'La navegación usa <Link> y <NavLink>. NavLink resuelve por sí mismo cuál es la '
     u'ruta activa, lo que sustituye al cálculo manual que hacía el NavbarComponent '
     u'de la Actividad 1:')
code_block([
    u'// src/components/Navbar.jsx',
    u'<NavLink to={path} end={path === ROUTES.CATALOG}',
    u'  className={({ isActive }) => [',
    u'    \'navlink\',',
    u'    isActive && \'navlink--active\',',
    u'    !isActive && cta && \'navlink--cta\',',
    u'  ].filter(Boolean).join(\' \')}>{label}</NavLink>',
], caption=u'El enlace activo se resalta y "Solicitar" se pinta como llamada a la '
           u'acción cuando no es la ruta actual.')

callout(u'Recargar en una ruta profunda',
        u'En desarrollo lo resuelve Vite. En producción, public/.htaccess viaja a '
        u'dist/ y reescribe cualquier ruta hacia index.html, de modo que recargar en '
        u'/simulador no devuelve 404 de Apache. El equivalente en Nginx es '
        u'try_files $uri $uri/ /index.html.',
        fill=WARN_BG, border=AMBER, icon='!')

page_break()

# ============================================================== 8. COMPONENTES
h1(u'8. Componentes y props')
para(u'Los once componentes son funciones que reciben props desestructuradas en la '
     u'firma y devuelven JSX. Uno por archivo. Ninguno guarda estado ni ejecuta '
     u'efectos: si un componente necesitara recordar algo, ese algo pertenece a la '
     u'página o a un hook.')
add_table(
    [u'Componente', u'Props principales', u'Responsabilidad'],
    [[u'Navbar', u'—', u'Barra superior; NavLink resuelve la ruta activa'],
     [u'Footer', u'variant', u'Pie con dos variantes: catalog y compact'],
     [u'CreditCard', u'product, variant',
      u'Tarjeta de producto; variante completa o compacta'],
     [u'SearchBar', u'value, onChange, label, placeholder',
      u'Entrada controlada de búsqueda'],
     [u'AmountRangeFilter', u'ranges, value, onChange',
      u'Desplegable de rangos de monto'],
     [u'SortSelect', u'options, value, onChange', u'Desplegable de ordenamiento'],
     [u'SimulatorForm', u'catalog, form, bounds, errors, onSelectProduct, '
      u'onAmountChange, onTermChange, onReset',
      u'Producto, monto y plazo, con los errores por campo'],
     [u'SimulationResult', u'simulation',
      u'Cuota, intereses, total y coste, ya formateados'],
     [u'AmortizationTable', u'simulation, isOpen, mode, onToggle, onModeChange',
      u'Tabla plegable en dos niveles de detalle'],
     [u'FormField', u'name, label, value, error, type, options, onChange, onBlur',
      u'input, select o textarea con etiqueta y error'],
     [u'Alert', u'variant, icon, children', u'Banda informativa reutilizable']],
    widths=[3.5, 5.6, 7.3])
cap(u'Los once componentes y sus props', kind='Tabla')

code_block([
    u'// src/components/CreditCard.jsx — props desestructuradas, sin estado',
    u'export function CreditCard({ product, variant = \'full\' }) {',
    u'  const { name, description, requirements, icon, themeClass,',
    u'          annualRateLabel, maxTermMonths, maxTermLabel,',
    u'          amountRangeLabel } = product;',
    u'',
    u'  const isCompact = variant === \'compact\';',
    u'  // …',
    u'  <div className="badge-amount">{amountRangeLabel}</div>',
    u'  {!isCompact && (',
    u'    <p className="product-card__requirements">',
    u'      <strong>Requisitos:</strong> {requirements}',
    u'    </p>',
    u'  )}',
    u'}',
], caption=u'Reutilización por variante: la misma tarjeta sirve al catálogo y al '
           u'simulador.')

callout(u'Por qué la tarjeta recibe un DTO y no una entidad',
        u'CreditProductDTO es un objeto plano y congelado, sin métodos. Si el '
        u'componente recibiera la entidad CreditProduct podría invocar reglas de '
        u'negocio desde la vista —por ejemplo decidir si un monto es admisible— y esa '
        u'decisión dejaría de estar en un solo sitio. Las cifras llegan ya '
        u'formateadas: la interfaz no formatea dinero.')

page_break()

# ============================================================== 9. ESTADO
h1(u'9. Manejo de estado con hooks')
h2(u'9.1 Hooks de React utilizados')
add_table(
    [u'Hook', u'Dónde y para qué'],
    [[u'useState', u'Texto de búsqueda, rango, criterio de orden, visibilidad de '
      u'tarjetas, producto/monto/plazo del simulador, apertura y modo de la tabla, '
      u'los 11 campos del formulario, campos tocados, envío en curso y radicado'],
     [u'useEffect', u'Cargar el catálogo, los rangos y los nombres de producto; '
      u'relanzar la búsqueda; recalcular la cuota; validar el borrador; fijar el '
      u'título del documento'],
     [u'useMemo', u'Ordenar y filtrar el catálogo sin recalcular en cada render; '
      u'construir una sola vez el objeto de dependencias'],
     [u'useCallback', u'Acciones estables (setValue, submit, reset, clearFilters) para '
      u'no re-renderizar los componentes hijos'],
     [u'useContext', u'Acceso a las dependencias inyectadas, dentro de useDependencies']],
    widths=[3.2, 13.2])
cap(u'Hooks de React y su uso en el proyecto', kind='Tabla')

h2(u'9.2 Los siete hooks propios')
add_table(
    [u'Hook', u'Qué encapsula', u'Caso de uso que invoca'],
    [[u'useDependencies', u'Acceso a los casos de uso inyectados', u'—'],
     [u'useDocumentTitle', u'Título del documento por ruta', u'—'],
     [u'useCreditProducts', u'Catálogo completo, carga y error',
      u'ListCreditProductsUseCase'],
     [u'useCreditSearch', u'Texto, rango, resultados y limpieza de filtros',
      u'SearchCreditProductsUseCase · GetAmountRangeFiltersUseCase'],
     [u'useProductSorting', u'Criterio de orden y lista ordenada', u'—'],
     [u'useSimulation', u'Producto, monto, plazo, simulación y errores',
      u'SimulateCreditUseCase'],
     [u'useApplicationForm', u'Los 11 campos, errores, envío y radicado',
      u'ValidateCreditApplicationDraftUseCase · SubmitCreditApplicationUseCase · '
      u'GetCreditProductNamesUseCase']],
    widths=[3.6, 6.4, 6.4])
cap(u'Hooks propios: el equivalente de los controladores de la Actividad 1',
    kind='Tabla')

para(u'La regla que ordena la capa: los hooks son los únicos que llaman a execute() y '
     u'los únicos que ven un Result. Fuera de un hook solo circulan datos planos, de '
     u'modo que ningún componente puede equivocarse interpretando un fallo.')

code_block([
    u'// src/hooks/useCreditProducts.js',
    u'useEffect(() => {',
    u'  let cancelled = false;              // React monta dos veces en modo estricto',
    u'',
    u'  async function load() {',
    u'    const result = await listCreditProducts.execute();',
    u'    if (cancelled) return;',
    u'',
    u'    if (result.isFailure) { setError(result.error); notifier.error(result.error); }',
    u'    else { setProducts(result.value.products); setTotal(result.value.total); }',
    u'    setIsLoading(false);',
    u'  }',
    u'',
    u'  load();',
    u'  return () => { cancelled = true; };   // limpieza: no escribir tras desmontar',
    u'}, [listCreditProducts, notifier]);',
], caption=u'Patrón común a todos los efectos asíncronos del proyecto.')

page_break()

# ============================================================== 10. BUSQUEDA
h1(u'10. Búsqueda en tiempo real y filtros dinámicos')
para(u'La búsqueda se aplica mientras el usuario escribe. No hay botón obligatorio: '
     u'cada pulsación actualiza el estado, y el cambio de estado relanza el caso de '
     u'uso. El criterio, en cambio, no se evalúa en la interfaz: lo resuelve el '
     u'dominio, con las mismas clases que usaría un backend.')
diagram([
    ('P', u'El usuario teclea «vehi»',
     u'SearchBar es una entrada controlada: notifica el valor con onChange'),
    ('P', u'useCreditSearch guarda el texto',
     u'setQuery dispara un nuevo render y, con él, el efecto de búsqueda'),
    ('A', u'SearchCreditProductsUseCase.execute()',
     u'Resuelve el rango por su índice y construye ProductSearchCriteria'),
    ('D', u'ProductSearchCriteria.matches()',
     u'CreditProduct.matchesName() y AmountRange.overlaps() deciden'),
    ('I', u'InMemoryCreditProductRepository',
     u'Aplica el criterio sobre el catálogo y devuelve entidades'),
    ('A', u'CreditProductMapper.toDTOList()',
     u'Convierte a DTOs planos y congelados, con los importes ya formateados'),
    ('P', u'La página pinta el resultado',
     u'Ordena, filtra por visibilidad y recorre con .map()'),
])
cap(u'Flujo completo de una pulsación en el buscador')

figure('captura-03-busqueda.jpg',
       u'Buscador, filtro por rango de monto, selector de orden y botón de limpiar '
       u'filtros, con el contador de resultados.', width_cm=15.5)

add_table(
    [u'Control', u'Qué hace', u'Dónde se decide'],
    [[u'Buscar por nombre', u'Filtra por nombre o descripción al teclear',
      u'Dominio — CreditProduct.matchesName()'],
     [u'Filtrar por rango de monto', u'5 rangos; el primero no acota',
      u'Dominio — AmountRange.overlaps()'],
     [u'Ordenar por', u'Nombre, tasa (asc/desc), monto máximo, plazo',
      u'Presentación — useProductSorting'],
     [u'Ocultar el producto simulado', u'No repite la tarjeta que ya se muestra arriba',
      u'Presentación — .filter() en la página'],
     [u'Limpiar filtros', u'Devuelve texto y rango a su valor inicial',
      u'Presentación — clearFilters']],
    widths=[4.2, 6.2, 6.0])
cap(u'Los cinco controles del catálogo y quién decide en cada uno', kind='Tabla')

callout(u'Por qué el orden se decide en la interfaz y el filtro por monto no',
        u'Que un producto de $5.000.000 pertenezca al rango «hasta $5.000.000» es una '
        u'afirmación sobre el negocio, y su respuesta debe ser la misma en cualquier '
        u'cliente. En qué orden se listan los resultados en pantalla, en cambio, es '
        u'una preferencia de presentación. Por eso lo primero vive en un value object '
        u'y lo segundo en un hook.')

page_break()

# ============================================================== 11. ARRAYS
h1(u'11. Manipulación de arrays: map, filter y sort')
para(u'Los tres métodos se usan encadenados sobre los DTOs recibidos. Como llegan '
     u'congelados desde la capa de aplicación, ordenar exige copiar: .sort() muta el '
     u'array que recibe.')
code_block([
    u'// src/hooks/useProductSorting.js — .sort() sobre una copia',
    u'const COMPARATORS = Object.freeze({',
    u'  name:        (a, b) => a.name.localeCompare(b.name, \'es\'),',
    u'  rateAsc:     (a, b) => a.annualRate - b.annualRate,',
    u'  rateDesc:    (a, b) => b.annualRate - a.annualRate,',
    u'  amountDesc:  (a, b) => (b.maxAmount ?? Infinity) - (a.maxAmount ?? Infinity),',
    u'  termDesc:    (a, b) => b.maxTermMonths - a.maxTermMonths,',
    u'});',
    u'',
    u'const sortedProducts = useMemo(',
    u'  () => [...products].sort(COMPARATORS[sortBy] ?? COMPARATORS.name),',
    u'  [products, sortBy],',
    u');',
], caption=u'Cinco criterios de orden, declarados como datos.')

code_block([
    u'// src/pages/SimulatorPage.jsx — .filter() y .map() con key unica',
    u'const visibleProducts = useMemo(',
    u'  () => sortedProducts.filter(',
    u'    (product) => !hideSimulated || String(product.id) !== String(form.productId),',
    u'  ),',
    u'  [sortedProducts, hideSimulated, form.productId],',
    u');',
    u'',
    u'{visibleProducts.map((product) => (',
    u'  <CreditCard key={product.id} product={product} variant="compact" />',
    u'))}',
], caption=u'La cadena completa: el dominio filtra por criterio, la interfaz ordena, '
           u'filtra por visibilidad y recorre.')

add_table(
    [u'Método', u'Dónde', u'Para qué'],
    [[u'.map()', u'CatalogPage, SimulatorPage, ApplicationPage, Navbar, '
      u'AmountRangeFilter, SortSelect, AmortizationTable, FormField',
      u'Recorrer listas; siempre con key única y estable (product.id, row.number, '
      u'field.name)'],
     [u'.filter()', u'SimulatorPage, Navbar',
      u'Descartar la tarjeta del producto simulado; componer clases CSS'],
     [u'.sort()', u'useProductSorting', u'Ordenar por los cinco criterios, sobre una copia'],
     [u'.reduce() / bucles', u'CreditSimulationService (dominio)',
      u'Construir la tabla de amortización y el resumen anual']],
    widths=[2.4, 6.6, 7.4])
cap(u'Uso de los métodos de array en el proyecto', kind='Tabla')

callout(u'La key no es un adorno',
        u'React usa la key para decidir qué nodos reutiliza entre renders. Con el '
        u'índice del array como key, reordenar la lista haría que React conservara el '
        u'nodo equivocado en cada posición: el foco, el scroll y las animaciones se '
        u'aplicarían a la tarjeta errónea. Por eso la key es el identificador del '
        u'dato, no su posición.',
        fill=WARN_BG, border=AMBER, icon='!')

page_break()

# ============================================================== 12. CUOTA
h1(u'12. Cálculo de la cuota mensual')
h2(u'12.1 La fórmula vive en el dominio')
para(u'La cuota se calcula con el sistema de amortización francés: cuota fija, '
     u'intereses sobre el saldo pendiente. La tasa nominal mensual se obtiene de la '
     u'tasa efectiva anual del producto elegido, y el único sitio del proyecto donde '
     u'aparece esta fórmula es un servicio de dominio.')
code_block([
    u'// domain/valueobjects/InterestRate.js — de tasa anual a mensual equivalente',
    u'i = (1 + tasaEfectivaAnual) ** (1 / 12) - 1',
    u'',
    u'// domain/services/CreditSimulationService.js — sistema frances',
    u'cuota = capital * i / (1 - (1 + i) ** (-n))',
    u'',
    u'// Para cada mes: interes = saldo * i ; capital = cuota - interes',
    u'// La ultima cuota absorbe el ajuste por redondeo al peso.',
], caption=u'Fórmula de la cuota. La interfaz no multiplica ni divide.')

h2(u'12.2 Recálculo automático')
para(u'No hay botón «Calcular». El hook del simulador guarda producto, monto y plazo, '
     u'y un efecto vuelve a invocar el caso de uso en cuanto cambia cualquiera de los '
     u'tres. Si el valor no es admisible, el resultado desaparece y se pintan los '
     u'errores por campo que produjo el dominio.')
code_block([
    u'// src/hooks/useSimulation.js',
    u'useEffect(() => {',
    u'  if (!form.productId) return undefined;',
    u'  let cancelled = false;',
    u'',
    u'  async function simulate() {',
    u'    const result = await simulateCredit.execute({',
    u'      productId: form.productId, amount: form.amount,',
    u'      termInMonths: form.termInMonths,',
    u'    });',
    u'    if (cancelled) return;',
    u'',
    u'    if (result.isFailure) { setSimulation(null); setErrors(result.fieldErrors); }',
    u'    else { setSimulation(result.value); setErrors({}); }',
    u'  }',
    u'',
    u'  simulate();',
    u'  return () => { cancelled = true; };',
    u'}, [form, simulateCredit]);',
], caption=u'Un solo efecto cubre los tres campos: la dependencia es el objeto form.')

h2(u'12.3 Resultado y formato en pesos')
figure('captura-02-simulador.jpg',
       u'Simulación de $1.000.000 a 12 meses al 18,5 % E.A.: cuota de $91.250, '
       u'$94.998 en intereses, $1.094.998 en total y un coste del 9,5 % del capital.',
       width_cm=15.5)

add_table(
    [u'Dato mostrado', u'Ejemplo', u'De dónde viene'],
    [[u'Cuota mensual', u'$ 91.250', u'AmortizationPlan.installment'],
     [u'Total en intereses', u'$ 94.998', u'AmortizationPlan.totalInterest'],
     [u'Total a pagar', u'$ 1.094.998', u'AmortizationPlan.totalPaid'],
     [u'Coste del crédito', u'9.5% del capital', u'SimulationMapper.interestRatioLabel'],
     [u'Última cuota', u'$ 91.248', u'Ajuste por redondeo al peso de las 12 cuotas'],
     [u'Formato COP', u'es-CO · COP', u'IntlMoneyFormatter, adaptador de IMoneyFormatter']],
    widths=[4.4, 3.8, 8.2])
cap(u'Cada cifra de la pantalla y su origen', kind='Tabla')

para(u'El formato en pesos no lo hace la interfaz: el dominio expone importes y un '
     u'puerto, IMoneyFormatter, declara la necesidad de convertirlos en texto. Su '
     u'adaptador usa Intl.NumberFormat con locale es-CO y moneda COP. Cambiar de '
     u'moneda es cambiar una línea en AppConfig.')

figure('chart-cuota-plazo.png',
       u'Efecto del plazo sobre la cuota y sobre el coste total, calculado con la '
       u'fórmula del dominio para un Crédito Vehículo de $20.000.000.', width_cm=15.5)

figure('chart-interes-capital.png',
       u'Reparto de las cuotas de ese crédito a 36 meses: el interés disminuye y el '
       u'abono a capital crece, mes a mes.', width_cm=14.5)

figure('chart-tasas.png',
       u'Tasa efectiva anual de los seis productos del catálogo. El simulador usa la '
       u'del producto seleccionado.', width_cm=15.0)

h2(u'12.4 Tabla de amortización')
para(u'La tabla va plegada por defecto y arranca en resumen anual: un crédito de '
     u'vivienda a 240 meses son 240 filas que nadie quiere de golpe. El detalle mes a '
     u'mes está a un clic. Ambos niveles llegan calculados en el DTO, con cada fila '
     u'cumpliendo la invariante cuota = interés + capital, verificada en el dominio.')
add_table(
    [u'Nivel', u'Columnas', u'Origen'],
    [[u'Resumen anual', u'Año · Cuotas · Pagado · Intereses · Capital · Saldo',
      u'AmortizationPlan.yearlySummary()'],
     [u'Detalle mensual', u'Cuota · Pago · Intereses · Capital · Saldo',
      u'AmortizationPlan.schedule (un Installment por mes)']],
    widths=[3.6, 7.4, 5.4])
cap(u'Los dos niveles de detalle de la tabla', kind='Tabla')

page_break()

# ============================================================== 13. FORMULARIO
h1(u'13. Formulario controlado y validaciones')
h2(u'13.1 Once campos, un solo estado')
para(u'El formulario es 100 % controlado: no existe ningún valor en el DOM que no '
     u'esté también en el estado de React. Los once campos viven en un único objeto, '
     u'y cada cambio pasa por el mismo manejador.')
add_table(
    [u'Sección', u'Campos', u'Value object que los valida'],
    [[u'Datos Personales', u'fullName · idNumber · email · phone', u'Applicant'],
     [u'Datos del Crédito', u'productName · amount · termInMonths · purpose',
      u'RequestedCredit'],
     [u'Datos Laborales', u'companyName · jobTitle · monthlyIncome', u'EmploymentInfo']],
    widths=[3.8, 8.2, 4.4])
cap(u'Las tres secciones del formulario y quién valida cada una', kind='Tabla')

code_block([
    u'// src/hooks/useApplicationForm.js',
    u'const setValue = useCallback((name, value) => {',
    u'  setValues((current) => ({ ...current, [name]: value }));',
    u'  setTouched((current) => ({ ...current, [name]: true }));',
    u'}, []);',
    u'',
    u'// Validacion en vivo con las reglas del dominio',
    u'useEffect(() => {',
    u'  let cancelled = false;',
    u'  (async () => {',
    u'    const result = await validateApplicationDraft.execute(values);',
    u'    if (!cancelled) setErrors(result.isSuccess ? {} : result.fieldErrors);',
    u'  })();',
    u'  return () => { cancelled = true; };',
    u'}, [values, validateApplicationDraft]);',
    u'',
    u'// Solo se muestra el error de un campo que el usuario ya toco',
    u'const errorFor = useCallback(',
    u'  (name) => (touched[name] ? (errors[name] ?? \'\') : \'\'),',
    u'  [errors, touched],',
    u');',
], caption=u'Estado, validación y visibilidad del error, en un solo hook.')

h2(u'13.2 Reglas de validación')
add_table(
    [u'Campo', u'Regla', u'Mensaje'],
    [[u'Nombre completo', u'Al menos nombre y apellido, mínimo 5 caracteres',
      u'Escribe el nombre y el apellido.'],
     [u'Cédula', u'Solo dígitos, entre 6 y 12',
      u'La cédula debe tener entre 6 y 12 dígitos.'],
     [u'Email', u'Formato usuario@dominio.tld', u'Ingresa un email válido.'],
     [u'Teléfono', u'10 dígitos, formato colombiano',
      u'El teléfono debe tener 10 dígitos.'],
     [u'Tipo de crédito', u'Debe existir en el catálogo',
      u'Selecciona un tipo de crédito.'],
     [u'Monto solicitado', u'Mayor que cero y dentro del rango del producto',
      u'Indica cuánto dinero necesitas. / El monto no está en el rango del producto.'],
     [u'Plazo', u'Entero mayor que cero y menor o igual al máximo del producto',
      u'El plazo debe ser un número entero de meses mayor que cero.'],
     [u'Destino del crédito', u'Descripción mínima', u'Describe el destino del crédito.'],
     [u'Empresa y cargo', u'Obligatorios', u'Completa los datos laborales.'],
     [u'Ingresos mensuales', u'Mayores que cero; la cuota no debe superar el 40 %',
      u'Con ese ingreso la cuota supera el 40 % permitido.']],
    widths=[3.4, 6.6, 6.4])
cap(u'Reglas de validación. Todas viven en el dominio, ninguna en el componente',
    kind='Tabla')

figure('captura-04-formulario.jpg',
       u'Validación en tiempo real: la cédula de tres dígitos y el correo incompleto '
       u'se señalan mientras se escribe, con el mensaje que produjo el dominio.',
       width_cm=15.5)

h2(u'13.3 Envío')
diagram([
    ('P', u'submit(event)', u'event.preventDefault(): no hay recarga de página'),
    ('P', u'Se marcan todos los campos como tocados',
     u'Al enviar se muestran todos los errores, no solo los de los campos visitados'),
    ('A', u'SubmitCreditApplicationUseCase.execute(values)',
     u'Traduce a value objects, construye la entidad y aplica la política'),
    ('D', u'CreditApplicationPolicy.assertAdmissible()',
     u'Monto y plazo admitidos por el producto; capacidad de pago'),
    ('D', u'application.submit()',
     u'Transición de estado válida; genera el número de radicado'),
    ('I', u'LocalStorageCreditApplicationRepository.save()',
     u'Persiste; degrada a memoria si el almacenamiento no está disponible'),
    ('P', u'Aviso y limpieza',
     u'ToastNotifier muestra el radicado; el formulario vuelve a su estado inicial'),
])
cap(u'Flujo de envío de una solicitud')

callout(u'Resultado real de la prueba',
        u'Con datos válidos, la aplicación radicó la solicitud con el número '
        u'CS-B946581E, mostró el aviso «Jeremy, tu solicitud quedó radicada con el '
        u'número CS-B946581E» y dejó el formulario vacío, sin errores en consola.')

page_break()

# ============================================================== 14. DI EN REACT
h1(u'14. Inyección de dependencias en React')
para(u'Sigue habiendo un solo Composition Root, y sigue siendo el único archivo del '
     u'proyecto donde se hace new de una clase concreta: config/dependencies.js. Lo '
     u'que cambió es cómo llegan esas dependencias a la interfaz.')
code_block([
    u'main.jsx                         <- construye el grafo y monta React',
    u' |- buildContainer()             <- config/dependencies.js: el unico `new`',
    u' |   \\- eagerResolveAll()        <- un contrato roto falla al cargar la pagina',
    u' \\- <DependenciesProvider>       <- traduce el contenedor a casos de uso',
    u'     \\- <BrowserRouter>',
    u'         \\- <App>                <- tabla de rutas',
    u'             \\- paginas -> hooks -> casos de uso',
], caption=u'Cadena de arranque de la aplicación.')

para(u'El detalle que preserva la regla de dependencia es que el proveedor no expone '
     u'el contenedor, sino un objeto congelado con los casos de uso ya resueltos:')
code_block([
    u'// src/context/DependenciesProvider.jsx',
    u'const services = useMemo(',
    u'  () => Object.freeze({',
    u'    listCreditProducts:   container.resolve(\'listCreditProductsUseCase\'),',
    u'    searchCreditProducts: container.resolve(\'searchCreditProductsUseCase\'),',
    u'    simulateCredit:       container.resolve(\'simulateCreditUseCase\'),',
    u'    // … y el resto de casos de uso',
    u'  }),',
    u'  [container],',
    u');',
], caption=u'Un componente no puede pedir un repositorio: no está en el objeto.')

para(u'Antes, la regla «la presentación nunca importa infraestructura» dependía de la '
     u'disciplina de quien escribía la vista. Ahora es imposible de romper por '
     u'construcción: lo único accesible desde un componente son casos de uso.')

add_table(
    [u'Bloque del contenedor', u'Registros', u'Contenido'],
    [[u'0. Valores de arranque', u'3',
      u'config · elemento de notificaciones · plazos del formulario'],
     [u'1. Adaptadores técnicos', u'5',
      u'ConsoleLogger · ToastNotifier · SystemClock · CryptoIdGenerator · '
      u'IntlMoneyFormatter'],
     [u'2. Persistencia', u'3',
      u'InMemoryCreditProductRepository · StaticAmountRangeProvider · '
      u'LocalStorageCreditApplicationRepository'],
     [u'3. Aplicación', u'9', u'2 mappers + 7 casos de uso'],
     [u'Total', u'20', u'Todos se construyen al arrancar con eagerResolveAll()']],
    widths=[4.6, 2.2, 9.6], align_center_cols=(1,))
cap(u'Grafo de dependencias registrado en config/dependencies.js', kind='Tabla')

h2(u'14.1 Contratos verificables en tiempo de ejecución')
para(u'JavaScript no tiene interface, así que los nueve puertos se declaran '
     u'explícitamente y se comprueban al construir el grafo.')
code_block([
    u'export const ICreditProductRepository = defineContract(',
    u'  \'ICreditProductRepository\',',
    u'  [\'findAll\', \'findById\', \'findByCriteria\', \'count\'],',
    u');',
    u'',
    u'// En el contenedor, al registrar el adaptador:',
    u'assertImplements(new InMemoryCreditProductRepository(), ICreditProductRepository);',
], caption=u'Si al adaptador le falta un método, la página no carga y el mensaje dice '
           u'exactamente cuál.')

page_break()

# ============================================================== 15. ARQUITECTURA
h1(u'15. La arquitectura que se conservó')
h2(u'15.1 Regla de dependencia')
para(u'Las flechas de dependencia apuntan siempre hacia el dominio. La presentación '
     u'conoce a la aplicación; la aplicación conoce al dominio; la infraestructura '
     u'conoce al dominio porque implementa sus puertos. El dominio no conoce a nadie.')
boxes_row([(u'presentación\nReact', BRAND), (u'application', VIOLET),
           (u'domain', NAVY)])
cap(u'Sentido de las dependencias. La infraestructura también apunta al dominio, '
    u'implementando sus puertos.')

para(u'Las tres reglas se comprueban con tres comandos, cambiando solo la ruta de la '
     u'presentación respecto a la Actividad 1:')
code_block([
    u'grep -rn "from \'../../\\(application\\|infrastructure\\|presentation\\|config\\)" src/domain/',
    u'grep -rn "from \'../../\\(infrastructure\\|presentation\\|config\\)"              src/application/',
    u'grep -rn "infrastructure" src/components src/pages src/hooks src/context src/App.jsx',
], caption=u'Las tres devuelven cero resultados.')

h2(u'15.2 Reparto de responsabilidades en la capa de React')
add_table(
    [u'Pieza', u'Puede', u'No puede'],
    [[u'components/', u'Recibir props y devolver JSX',
      u'Tener estado, efectos ni llamar a casos de uso'],
     [u'pages/', u'Componer, decidir qué se muestra y en qué orden',
      u'Consultar datos por su cuenta'],
     [u'hooks/', u'Guardar estado de UI, invocar casos de uso, desenvolver Result',
      u'Contener reglas de negocio'],
     [u'context/', u'Entregar los casos de uso al árbol de componentes',
      u'Exponer el contenedor o un adaptador']],
    widths=[3.0, 7.0, 6.4])
cap(u'Qué puede y qué no puede cada pieza de la presentación', kind='Tabla')

h2(u'15.3 Principios SOLID en la nueva capa')
add_table(
    [u'Principio', u'Dónde se ve en la Actividad 2'],
    [[u'Responsabilidad única', u'Un componente pinta, un hook coordina, un caso de '
      u'uso orquesta, una entidad decide. FormField solo sabe de un campo'],
     [u'Abierto/cerrado', u'Añadir un producto es añadir una entrada en '
      u'creditsData.js: ningún componente cambia'],
     [u'Sustitución de Liskov', u'Cualquier adaptador de ICreditProductRepository '
      u'sirve; los métodos son async aunque no haya E/S'],
     [u'Segregación de interfaces', u'Nueve puertos pequeños, de 2 a 5 métodos, en vez '
      u'de una interfaz general'],
     [u'Inversión de dependencias', u'Los componentes reciben casos de uso por '
      u'contexto; el dominio depende de puertos, no de adaptadores']],
    widths=[3.6, 12.8])
cap(u'Los cinco principios en la capa de presentación de React', kind='Tabla')

page_break()

# ============================================================== 16. ESTILOS
h1(u'16. Estilos y diseño responsive')
para(u'Los siete archivos CSS de la Actividad 1 se reutilizaron sin modificar ni una '
     u'regla. Se importan una sola vez en main.jsx, en orden, para que Vite los '
     u'empaquete respetando la cascada.')
code_block([
    u'// src/main.jsx',
    u'import \'../assets/css/01-reset.css\';',
    u'import \'../assets/css/02-tokens.css\';',
    u'import \'../assets/css/03-base.css\';',
    u'import \'../assets/css/04-layout.css\';',
    u'import \'../assets/css/05-components.css\';',
    u'import \'../assets/css/06-pages.css\';',
    u'import \'../assets/css/07-responsive.css\';',
], caption=u'Cascada explícita: reset → tokens → base → layout → componentes → '
           u'páginas → responsive.')

add_table(
    [u'Archivo', u'Contenido'],
    [[u'01-reset.css', u'Normalización de márgenes, tipografía y caja'],
     [u'02-tokens.css', u'Variables de color, espaciado, sombras y las 6 paletas de tema'],
     [u'03-base.css', u'Tipografía, enlaces, títulos y controles de formulario'],
     [u'04-layout.css', u'Contenedores, rejillas y estructura de página'],
     [u'05-components.css', u'Tarjetas, botones, campos, avisos, barra y pie'],
     [u'06-pages.css', u'Hero, panel del simulador, resultado, tabla y formulario'],
     [u'07-responsive.css', u'Puntos de quiebre móvil, tableta y escritorio']],
    widths=[4.2, 12.2])
cap(u'Los siete archivos de estilos, reutilizados de la Actividad 1', kind='Tabla')

para(u'Que el diseño se conserve sin tocar el CSS no es casualidad: los componentes de '
     u'React reproducen las mismas clases y la misma estructura de marcado que las '
     u'plantillas de la Actividad 1. Fue una decisión deliberada, y evitó rehacer el '
     u'trabajo de maquetación y de responsive ya evaluado.')

figure('captura-01-catalogo.jpg',
       u'Catálogo en React con el mismo diseño de la Actividad 1: hero, rejilla de '
       u'productos y tarjetas con tema por producto.', width_cm=15.5)

page_break()

# ============================================================== 17. VERIFICACION
h1(u'17. Verificación')
h2(u'17.1 Comprobaciones automáticas')
add_table(
    [u'Comprobación', u'Comando', u'Resultado'],
    [[u'Reglas de negocio intactas', u'node tests/01-domain-application.mjs',
      u'85 aserciones · TODO OK'],
     [u'Regla de dependencia (dominio)', u'grep sobre src/domain/', u'0 resultados'],
     [u'Regla de dependencia (aplicación)', u'grep sobre src/application/',
      u'0 resultados'],
     [u'Presentación sin infraestructura',
      u'grep sobre components, pages, hooks, context, App.jsx', u'0 resultados'],
     [u'Compilación', u'npm run build', u'Sin errores ni advertencias'],
     [u'Contratos', u'container.eagerResolveAll() al arrancar',
      u'Las 20 dependencias se construyen sin violaciones']],
    widths=[4.6, 6.4, 5.4])
cap(u'Comprobaciones automáticas y su resultado', kind='Tabla')

para(u'La suite 01 es la prueba central de la tesis del documento: cubre entidades, '
     u'value objects, servicios de dominio, casos de uso, Result y DTOs, y pasó sin '
     u'modificar una sola línea después de reescribir toda la interfaz. Las suites 02 '
     u'y 03 de la Actividad 1 se retiraron porque probaban el render de plantillas de '
     u'cadena y el arranque con jsdom, dos cosas que ya no existen; su papel lo asume '
     u'npm run build junto con la comprobación manual en el navegador.')

h2(u'17.2 Comprobaciones en el navegador')
add_table(
    [u'Escenario', u'Resultado observado'],
    [[u'Navegación entre las cuatro rutas', u'Sin recarga; el enlace activo se resalta'],
     [u'Recarga directa en /simulador y /solicitar', u'La página se sirve correctamente'],
     [u'Búsqueda «vehi»',
      u'Filtra al teclear: «Mostrando 1 de 6 productos», solo Crédito Vehículo'],
     [u'Filtro por rango y limpiar filtros', u'Resultados y contador coherentes'],
     [u'Ordenar por nombre',
      u'de Libranza · Educativo · Empresarial · Libre Inversión · Vehículo · Vivienda'],
     [u'Cambio de monto o plazo', u'La cuota se recalcula sin pulsar ningún botón'],
     [u'Tabla de amortización', u'Se pliega y alterna entre resumen anual y mes a mes'],
     [u'Cédula «123» y correo «juan@»',
      u'Mensajes por campo mientras se escribe, solo en los campos tocados'],
     [u'Envío con datos válidos',
      u'Radicado CS-B946581E, aviso al usuario y formulario limpio'],
     [u'Ruta inexistente', u'Pantalla 404 con la ruta solicitada y enlace a inicio'],
     [u'Consola del navegador', u'Sin errores ni advertencias']],
    widths=[5.6, 10.8])
cap(u'Pruebas manuales realizadas en el navegador', kind='Tabla')

h2(u'17.3 Diseño responsive')
para(u'Verificado en los tres tamaños con las hojas de estilo heredadas: rejilla de '
     u'productos de tres columnas en escritorio, dos en tableta y una en móvil; barra '
     u'de navegación y formulario apilados en pantallas estrechas; tabla de '
     u'amortización con desplazamiento horizontal propio.')
field_placeholder(u'[ Opcional: pegue aquí capturas en móvil y tableta ]',
                  width_cm=15.0, height_cm=4.5)

page_break()

# ============================================================== 18. RUBRICA
h1(u'18. Cumplimiento de la rúbrica')
add_table(
    [u'Criterio', u'Pts', u'Evidencia en el proyecto', u'Auto'],
    [[u'Configuración y estructura de React', u'20',
      u'Vite + React Router configurados; carpetas components/, pages/, data/, '
      u'hooks/, context/; archivos nombrados por artefacto (§6, §7)', u''],
     [u'Componentes y props', u'15',
      u'11 componentes funcionales, uno por archivo, props desestructuradas en la '
      u'firma; dos variantes reutilizables en CreditCard (§8)', u''],
     [u'Manejo de estado (useState)', u'10',
      u'Estados nombrados e inicializados en 7 hooks propios; actualización '
      u'inmutable con setters y forma funcional (§9)', u''],
     [u'Búsqueda y filtros dinámicos', u'15',
      u'Búsqueda al teclear, filtro por rango de monto, orden y botón de limpiar '
      u'filtros; contador de resultados (§10)', u''],
     [u'Formulario controlado', u'10',
      u'11 campos controlados, validación en vivo de correo, cédula y montos, '
      u'mensajes por campo, preventDefault y limpieza tras el envío (§13)', u''],
     [u'Manipulación de arrays', u'10',
      u'.map() con key única, .filter() y .sort() encadenados sobre copias (§11)', u''],
     [u'Cálculo de cuota mensual', u'10',
      u'Recálculo automático con la tasa del producto, formato COP y tabla de '
      u'amortización (§12)', u''],
     [u'Sustentación sincrónica', u'10',
      u'Guion preparado en docs/iudigital_doc/EV2/guion-sustentacion-EV2.md', u''],
     [u'Total', u'100', u'', u'']],
    widths=[4.0, 1.2, 9.4, 1.8], align_center_cols=(1, 3))
cap(u'Criterios de la rúbrica y evidencia. La última columna queda para la '
    u'autoevaluación', kind='Tabla')

h2(u'18.1 Observaciones del estudiante')
field_placeholder(u'[ Escriba aquí las observaciones que quiera añadir a la entrega ]',
                  width_cm=15.0, height_cm=3.0)

page_break()

# ============================================================== 19. ENTREGABLES
h1(u'19. Entregables')
h2(u'19.1 Repositorio')
field_placeholder(u'[ URL del repositorio Git ]', width_cm=13.0)
add_table(
    [u'Elemento', u'Estado'],
    [[u'Código fuente completo en src/', u'Incluido'],
     [u'package.json con dependencias', u'Incluido, con package-lock.json versionado'],
     [u'.gitignore configurado', u'node_modules/ y dist/ excluidos'],
     [u'README.md actualizado',
      u'Integrantes, descripción, tecnologías, instalación y capturas'],
     [u'Capturas de pantalla', u'4, en docs/capturas/'],
     [u'Documentación de diseño', u'21 documentos en docs/, índice en docs/master.md'],
     [u'Actividad 1 preservada', u'Tag ev1-entrega, en la misma historia del repositorio']],
    widths=[6.0, 10.4])
cap(u'Contenido del repositorio', kind='Tabla')

h2(u'19.2 Historial de commits')
para(u'La entrega se desarrolló en la rama ev2-react, con la Actividad 1 congelada en '
     u'el tag ev1-entrega para que siga siendo reproducible.')
add_table(
    [u'#', u'Commit', u'Qué aporta'],
    [[u'1', u'chore: configura React con Vite y prepara el punto de montaje',
      u'package.json, vite.config.js, index.html de Vite'],
     [u'2', u'refactor: extrae el catálogo a src/data/creditsData.js',
      u'Carpeta data/ con el archivo de datos'],
     [u'3', u'feat: sustituye la presentación vanilla por React con React Router',
      u'Proveedor de dependencias, rutas, Navbar, Footer, CreditCard, catálogo'],
     [u'4', u'feat: implementa la búsqueda en tiempo real y el filtro por rango',
      u'SearchBar, AmountRangeFilter, useCreditSearch, limpiar filtros'],
     [u'5', u'feat: añade ordenamiento del catálogo con .sort()',
      u'SortSelect y useProductSorting'],
     [u'6', u'feat: calcula la cuota mensual y muestra la tabla de amortización',
      u'useSimulation, SimulatorForm, SimulationResult, AmortizationTable'],
     [u'7', u'feat: implementa el formulario de solicitud controlado con validaciones',
      u'useApplicationForm, FormField, caso de uso de validación'],
     [u'8', u'docs: documenta la migración a React y actualiza el patrón',
      u'Documento 21 y actualización de master.md y de los docs afectados'],
     [u'9', u'docs: actualiza el README para la Actividad 2 y añade capturas',
      u'README con integrantes, tecnologías, instalación y 4 capturas']],
    widths=[0.9, 7.5, 8.0], align_center_cols=(0,))
cap(u'Los nueve commits de la rama ev2-react', kind='Tabla')

field_placeholder(u'[ Opcional: pegue aquí la captura del historial de commits ]',
                  width_cm=15.0, height_cm=4.0)

page_break()

# ============================================================== 20. CONCLUSIONES
h1(u'20. Conclusiones')
for t in (u'La transformación del diseño estático en una aplicación interactiva con '
          u'React se completó cubriendo los siete criterios técnicos de la rúbrica: '
          u'configuración con Vite y React Router, componentes con props, estado con '
          u'hooks, búsqueda y filtros dinámicos, formulario controlado, manipulación '
          u'de arrays y cálculo de la cuota mensual.',
          u'La arquitectura hexagonal convirtió un cambio de framework en un cambio '
          u'de adaptador. La capa de dominio —26 archivos con todas las reglas del '
          u'negocio— no se modificó, y la suite que la cubre siguió pasando sin '
          u'ajustes después de reescribir la interfaz completa. Es la evidencia '
          u'empírica de una afirmación que en la Actividad 1 era todavía teórica.',
          u'Los hooks propios resultaron el equivalente natural de los controladores: '
          u'guardan el estado de la interfaz, invocan casos de uso y desenvuelven el '
          u'Result. Concentrar ahí esa responsabilidad dejó a los componentes libres '
          u'de lógica y permitió que la regla «la presentación no importa '
          u'infraestructura» dejara de depender de la disciplina del programador.',
          u'Validar en tiempo real reutilizando los value objects del dominio evitó el '
          u'problema más común de los formularios controlados: dos validaciones '
          u'distintas, una en el cliente y otra en el servidor, que se desincronizan. '
          u'El caso de uso de borrador cuesta doce líneas y elimina esa duplicación.',
          u'Reutilizar los siete archivos CSS sin modificarlos confirmó que el diseño '
          u'y la estructura del marcado son independientes de la tecnología que los '
          u'genera. El coste de la migración fue proporcional al tamaño de la '
          u'presentación, no al del proyecto.',
          u'Queda como extensión natural sustituir el datasource estático por una API '
          u'HTTP: es un adaptador nuevo y una línea en el contenedor, sin tocar el '
          u'dominio, los casos de uso ni un solo componente.'):
    bullet(t)

# ============================================================== 21. REFERENCIAS
h1(u'21. Referencias')
for ref in (
    u'Meta Open Source. React — Documentación oficial. react.dev',
    u'Meta Open Source. Reutilizar lógica con Hooks personalizados. react.dev/learn',
    u'Remix Software. React Router — Documentación oficial. reactrouter.com',
    u'Vite. Guía de inicio y configuración. vite.dev',
    u'MDN Web Docs. Array.prototype.map, filter y sort. developer.mozilla.org',
    u'MDN Web Docs. Intl.NumberFormat. developer.mozilla.org',
    u'Cockburn, A. Hexagonal Architecture (Ports and Adapters), 2005.',
    u'Martin, R. C. Clean Architecture: A Craftsman’s Guide to Software Structure and '
    u'Design. Prentice Hall, 2017.',
    u'Evans, E. Domain-Driven Design: Tackling Complexity in the Heart of Software. '
    u'Addison-Wesley, 2003.',
    u'Documentación interna del proyecto: docs/master.md (índice de los 21 documentos '
    u'de diseño) y docs/21-migracion-a-react-ev2.md (delta de esta actividad).',
):
    bullet(ref)

para(u'Nota sobre el origen del contenido funcional: la aplicación es una '
     u'reconstrucción del sitio de referencia sweet-smart-credit-path.base44.app. Se '
     u'replicaron sus textos, productos, colores y puntos de quiebre; la arquitectura, '
     u'el enrutado, las validaciones, el simulador y las pruebas son desarrollo '
     u'propio.', size=9, italic=True, color='6B7280')

doc.save(OUT_DOCX)
print(u'Documento generado: %s' % OUT_DOCX)
print(u'Figuras: %d   Tablas numeradas: %d' % (FIG_N[0], TAB_N[0]))
